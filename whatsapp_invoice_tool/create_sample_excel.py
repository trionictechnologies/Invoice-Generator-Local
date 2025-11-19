"""
Script to create a sample Excel file for testing the WhatsApp Invoice Tool.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime

def create_sample_excel():
    """Create a sample Excel file with invoice data."""
    
    # Create workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Invoices"
    
    # Headers
    headers = [
        "InvoiceNo", 
        "CustomerName", 
        "PhoneNumber", 
        "Amount", 
        "InvoiceDate", 
        "Status", 
        "PaymentReceived"
    ]
    
    # Style headers
    header_fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Sample data
    sample_data = [
        ["INV-2024-001", "Rajesh Kumar", "919876543210", 15000.00, "2024-01-15", "Pending", ""],
        ["INV-2024-002", "Priya Sharma", "919876543211", 25000.50, "2024-01-16", "Pending", ""],
        ["INV-2024-003", "Amit Patel", "919876543212", 8500.00, "2024-01-17", "Pending", ""],
        ["INV-2024-004", "Sunita Reddy", "919876543213", 32000.00, "2024-01-18", "Pending", ""],
        ["INV-2024-005", "Vikram Singh", "919876543214", 12750.75, "2024-01-19", "Pending", ""],
    ]
    
    # Add data rows
    for row_num, row_data in enumerate(sample_data, 2):
        for col_num, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_num, column=col_num, value=value)
            cell.alignment = Alignment(horizontal="left", vertical="center")
    
    # Adjust column widths
    column_widths = [15, 20, 15, 12, 15, 12, 18]
    for col_num, width in enumerate(column_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_num)].width = width
    
    # Save workbook
    filename = "sample_invoices.xlsx"
    wb.save(filename)
    print(f"✓ Sample Excel file created: {filename}")
    print(f"  - Contains {len(sample_data)} sample invoices")
    print(f"  - All marked as 'Pending' status")
    print(f"\nNote: Replace phone numbers with actual numbers before testing!")

if __name__ == "__main__":
    create_sample_excel()
