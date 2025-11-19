"""
Excel Manager for WhatsApp Invoice Tool.
Handles reading and writing invoice data from Excel files with multi-sheet support.
Supports GST invoice format with column mapping.
"""
import openpyxl
from typing import List, Dict, Optional
from pathlib import Path
from datetime import datetime
from .logger import logger


class ExcelManager:
    """Manages Excel file operations for invoice data with multi-sheet support."""
    
    def __init__(self, file_path: str, config: Optional[Dict] = None):
        """
        Initialize Excel Manager.
        
        Args:
            file_path: Path to the Excel file
            config: Optional configuration dict
        """
        self.file_path = Path(file_path)
        
        if not self.file_path.exists():
            raise FileNotFoundError(f"Excel file not found: {file_path}")
        
        self.config = config or {}
        self.column_mapping = {}
        self.data_sheet_name = self.config.get("data_sheet_name", "InvoiceData")
        self.mapping_sheet_name = self.config.get("mapping_sheet_name", "ColumnMapping")
        
        # Load column mapping from Excel or use defaults
        self._load_column_mapping()
    
    def _load_column_mapping(self):
        """Load column mapping from Excel file or use default mapping."""
        try:
            workbook = openpyxl.load_workbook(self.file_path, data_only=True)
            
            # Check if mapping sheet exists
            if self.mapping_sheet_name in workbook.sheetnames:
                mapping_sheet = workbook[self.mapping_sheet_name]
                logger.info(f"Loading column mapping from '{self.mapping_sheet_name}' sheet")
                
                # Read mapping: Column 1 = Field Name, Column 2 = Excel Column Letter
                for row in range(2, mapping_sheet.max_row + 1):
                    field_name = mapping_sheet.cell(row=row, column=1).value
                    column_letter = mapping_sheet.cell(row=row, column=2).value
                    
                    if field_name and column_letter:
                        # Convert column letter to index (A=1, B=2, etc.)
                        col_index = openpyxl.utils.column_index_from_string(str(column_letter).strip().upper())
                        self.column_mapping[str(field_name).strip()] = col_index
                
                logger.info(f"Loaded {len(self.column_mapping)} column mappings")
            else:
                logger.info(f"No mapping sheet found, using default column order")
                # Use default sequential mapping if no mapping sheet
                self._set_default_mapping()
            
            workbook.close()
            
        except Exception as e:
            logger.warning(f"Error loading column mapping: {str(e)}. Using defaults.")
            self._set_default_mapping()
    
    def _set_default_mapping(self):
        """Set default column mapping (sequential from A onwards)."""
        # Default GST invoice field mapping
        default_fields = [
            "Logo", "SupplierName", "SupplierAddress", "SupplierGSTIN",
            "BuyerName", "BuyerAddress", "BuyerGSTIN", "InvoiceNumber",
            "InvoiceDate", "Description", "SAC_HSN", "Total",
            "CGST", "SGST", "IGST", "GrandTotal", "AmountInWords",
            "BankDetails", "TermsOfDelivery", "OtherTerms", "SignatureField",
            "PhoneNumber", "Status", "PaymentReceived"
        ]
        
        for idx, field in enumerate(default_fields, start=1):
            self.column_mapping[field] = idx
    
    def _get_field_value(self, row, field_name: str):
        """
        Get value for a specific field from the row.
        
        Args:
            row: openpyxl row object
            field_name: Name of the field
            
        Returns:
            Cell value or None
        """
        col_index = self.column_mapping.get(field_name)
        if col_index is None:
            return None
        
        try:
            return row[col_index - 1].value
        except (IndexError, TypeError):
            return None
    
    def load_invoices(self) -> List[Dict]:
        """
        Load all invoice rows from Excel file.
        
        Returns:
            List of invoice dictionaries with all GST invoice fields
        """
        try:
            workbook = openpyxl.load_workbook(self.file_path, data_only=True)
            
            # Try to get the data sheet, fallback to active sheet
            if self.data_sheet_name in workbook.sheetnames:
                sheet = workbook[self.data_sheet_name]
                logger.info(f"Reading data from '{self.data_sheet_name}' sheet")
            else:
                sheet = workbook.active
                logger.info(f"Reading data from active sheet: '{sheet.title}'")
            
            invoices = []
            
            # Skip header row (row 1), start from row 2
            for row_num in range(2, sheet.max_row + 1):
                row = sheet[row_num]
                
                # Read invoice number to check if row is valid
                invoice_no = self._get_field_value(row, "InvoiceNumber")
                if not invoice_no:
                    continue
                
                # Read all fields
                phone_number = self._get_field_value(row, "PhoneNumber")
                invoice_date = self._get_field_value(row, "InvoiceDate")
                status = self._get_field_value(row, "Status")
                payment_received = self._get_field_value(row, "PaymentReceived")
                
                # Clean phone number
                if phone_number:
                    phone_number = str(phone_number).replace(" ", "").replace("-", "").replace("+", "")
                
                # Convert date
                if isinstance(invoice_date, datetime):
                    invoice_date_str = invoice_date.strftime("%d-%b-%Y")
                elif invoice_date:
                    invoice_date_str = str(invoice_date)
                else:
                    invoice_date_str = ""
                
                # Read all GST invoice fields
                invoice_data = {
                    # Supplier Info
                    "Logo": self._get_field_value(row, "Logo") or "",
                    "SupplierName": self._get_field_value(row, "SupplierName") or "",
                    "SupplierAddress": self._get_field_value(row, "SupplierAddress") or "",
                    "SupplierGSTIN": self._get_field_value(row, "SupplierGSTIN") or "",
                    
                    # Buyer Info
                    "BuyerName": self._get_field_value(row, "BuyerName") or "",
                    "BuyerAddress": self._get_field_value(row, "BuyerAddress") or "",
                    "BuyerGSTIN": self._get_field_value(row, "BuyerGSTIN") or "",
                    
                    # Invoice Details
                    "InvoiceNumber": str(invoice_no),
                    "InvoiceDate": invoice_date_str,
                    "Description": self._get_field_value(row, "Description") or "Services Rendered",
                    "SAC_HSN": self._get_field_value(row, "SAC_HSN") or "",
                    
                    # Amounts
                    "Total": float(self._get_field_value(row, "Total") or 0),
                    "CGST": float(self._get_field_value(row, "CGST") or 0),
                    "SGST": float(self._get_field_value(row, "SGST") or 0),
                    "IGST": float(self._get_field_value(row, "IGST") or 0),
                    "GrandTotal": float(self._get_field_value(row, "GrandTotal") or 0),
                    "AmountInWords": self._get_field_value(row, "AmountInWords") or "",
                    
                    # Additional Info
                    "BankDetails": self._get_field_value(row, "BankDetails") or "",
                    "TermsOfDelivery": self._get_field_value(row, "TermsOfDelivery") or "",
                    "OtherTerms": self._get_field_value(row, "OtherTerms") or "",
                    "SignatureField": self._get_field_value(row, "SignatureField") or "",
                    
                    # WhatsApp & Status
                    "PhoneNumber": phone_number or "",
                    "Status": str(status) if status else "Pending",
                    "PaymentReceived": str(payment_received) if payment_received else "",
                    
                    # Internal
                    "RowNumber": row_num
                }
                
                # Backward compatibility: also set CustomerName and Amount
                invoice_data["CustomerName"] = invoice_data["BuyerName"]
                invoice_data["Amount"] = invoice_data["GrandTotal"]
                invoice_data["InvoiceNo"] = invoice_data["InvoiceNumber"]
                
                invoices.append(invoice_data)
            
            workbook.close()
            logger.info(f"Loaded {len(invoices)} invoices from {self.file_path}")
            return invoices
            
        except Exception as e:
            logger.error(f"Error loading invoices from Excel: {str(e)}")
            raise
    
    def update_status(self, invoice_no: str, status: str) -> bool:
        """
        Update the status column for a specific invoice.
        
        Args:
            invoice_no: Invoice number to update
            status: New status value (e.g., "Sent", "Error:...")
            
        Returns:
            True if successful, False otherwise
        """
        try:
            workbook = openpyxl.load_workbook(self.file_path)
            
            # Get data sheet
            if self.data_sheet_name in workbook.sheetnames:
                sheet = workbook[self.data_sheet_name]
            else:
                sheet = workbook.active
            
            status_col = self.column_mapping.get("Status")
            invoice_col = self.column_mapping.get("InvoiceNumber")
            
            if not status_col or not invoice_col:
                logger.error("Status or InvoiceNumber column not mapped")
                workbook.close()
                return False
            
            # Find the row with matching invoice number
            for row_num in range(2, sheet.max_row + 1):
                row = sheet[row_num]
                current_invoice_no = str(self._get_field_value(row, "InvoiceNumber"))
                
                if current_invoice_no == str(invoice_no):
                    # Update status column
                    sheet.cell(row=row_num, column=status_col, value=status)
                    workbook.save(self.file_path)
                    workbook.close()
                    logger.info(f"Updated status for invoice {invoice_no} to '{status}'")
                    return True
            
            workbook.close()
            logger.warning(f"Invoice {invoice_no} not found in Excel")
            return False
            
        except Exception as e:
            logger.error(f"Error updating status for invoice {invoice_no}: {str(e)}")
            return False
    
    def update_payment_received(self, invoice_no: str, value: str) -> bool:
        """
        Update the PaymentReceived column for a specific invoice.
        
        Args:
            invoice_no: Invoice number to update
            value: Payment received value (timestamp or "Yes")
            
        Returns:
            True if successful, False otherwise
        """
        try:
            workbook = openpyxl.load_workbook(self.file_path)
            
            # Get data sheet
            if self.data_sheet_name in workbook.sheetnames:
                sheet = workbook[self.data_sheet_name]
            else:
                sheet = workbook.active
            
            payment_col = self.column_mapping.get("PaymentReceived")
            invoice_col = self.column_mapping.get("InvoiceNumber")
            
            if not payment_col or not invoice_col:
                logger.error("PaymentReceived or InvoiceNumber column not mapped")
                workbook.close()
                return False
            
            # Find the row with matching invoice number
            for row_num in range(2, sheet.max_row + 1):
                row = sheet[row_num]
                current_invoice_no = str(self._get_field_value(row, "InvoiceNumber"))
                
                if current_invoice_no == str(invoice_no):
                    # Update payment received column
                    sheet.cell(row=row_num, column=payment_col, value=value)
                    workbook.save(self.file_path)
                    workbook.close()
                    logger.info(f"Updated payment received for invoice {invoice_no} to '{value}'")
                    return True
            
            workbook.close()
            logger.warning(f"Invoice {invoice_no} not found in Excel")
            return False
            
        except Exception as e:
            logger.error(f"Error updating payment received for invoice {invoice_no}: {str(e)}")
            return False
